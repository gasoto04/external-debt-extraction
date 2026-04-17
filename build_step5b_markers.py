"""Step 5b: fill markers into empty dsa.pdf cells."""
import os, re, warnings
from collections import defaultdict
warnings.filterwarnings('ignore')
from openpyxl import load_workbook

COMPILED_DIR = '/home/gabos/projects/external-debt-extraction/after-2018/_compiled'
PAIRS_PATH = '/home/gabos/projects/external-debt-extraction/dsa_parent_pairs.xlsx'

NAME_MAP = {
    'Congo DR': 'Congo, Dem. Rep',
    'Congo Republic of': 'Congo, Rep',
    "Cote d'Ivoire": "Côte d'Ivoire",
    'Gambia': 'Gambia, The',
    'Micronesia': 'Micronesia, Federated States of',
    'St. Vincent': 'St. Vincent and the Grenadines',
    'Yemen': 'Yemen, Rep',
}
TYPO_FIX = {'Afghanisatn': 'Afghanistan', "Cote d'lvoire": "Cote d'Ivoire", 'Tazania': 'Tanzania'}

# Load pairs
wb = load_workbook(PAIRS_PATH, read_only=True)
ws = wb.active
pairs = defaultdict(lambda: defaultdict(list))
for row in ws.iter_rows(min_row=2, values_only=True):
    fn, country, _, year, month = row[0], row[1], row[2], row[3], row[4]
    if fn and country and year and month:
        try: pairs[country][int(year)].append((int(month), fn))
        except (TypeError, ValueError): pass
wb.close()
for c in pairs:
    for y in pairs[c]: pairs[c][y].sort()

SHEET_RE = re.compile(r'^([A-Z]{3})_(EBS|SM)\.(\d{2})\.(\d+)')

# Recompute what marker each empty sheet should get
def marker_for(cname_pairs, year, group_ours):
    """group_ours is list of (doc, sheet) sorted ascending. Returns dict {sheet -> marker}."""
    pair_list = pairs.get(cname_pairs, {}).get(year, [])
    out = {}
    n_pairs = len(pair_list)
    if n_pairs == 0:
        marker = 'no_pdf_2024' if year == 2024 else None  # None = leave empty
        for doc, s in group_ours: out[s] = ('EMPTY_NOPAIR', marker)
        return out
    # first min(n) get matched; overflow sheets get MORE_EXCEL marker
    n = min(len(group_ours), n_pairs)
    for i in range(n):
        out[group_ours[i][1]] = ('MATCHED', None)
    for i in range(n, len(group_ours)):
        out[group_ours[i][1]] = ('MORE', 'more_excel_than_pdf')
    return out

# Gather entries
entries = defaultdict(list)  # (cf, cname_pairs, year) -> [(doc, sheet)]
for cf in sorted(os.listdir(COMPILED_DIR)):
    if not cf.endswith('.xlsx') or cf.startswith('_') or cf.startswith('~$'): continue
    folder = os.path.splitext(cf)[0]
    cname_raw = folder.rsplit('_', 1)[0].strip()
    cname = TYPO_FIX.get(cname_raw, cname_raw)
    cname_pairs = NAME_MAP.get(cname, cname)
    wb = load_workbook(os.path.join(COMPILED_DIR, cf), read_only=True)
    for s in wb.sheetnames:
        if s == 'EMPTY': continue
        m = SHEET_RE.match(s)
        if not m: continue
        yy = int(m.group(3))
        year = 2000 + yy if yy < 50 else 1900 + yy
        entries[(cf, cname_pairs, year)].append((int(m.group(4)), s))
    wb.close()

# Compute markers per sheet
sheet_markers = defaultdict(dict)  # cf -> {sheet -> marker_value_or_None}
still_empty_non2024 = []
for (cf, cname_pairs, year), ours in entries.items():
    ours.sort()
    info = marker_for(cname_pairs, year, ours)
    for s, (kind, marker) in info.items():
        if kind == 'MATCHED': continue  # leave as is
        sheet_markers[cf][s] = marker
        if marker is None:
            still_empty_non2024.append((cf, s, cname_pairs, year))

# Apply markers (only write if current A2 is empty/None)
updated = 0
for cf, smap in sheet_markers.items():
    path = os.path.join(COMPILED_DIR, cf)
    wb = load_workbook(path)
    changed = False
    for sname, marker in smap.items():
        if marker is None: continue
        ws = wb[sname]
        if ws['A2'].value is None:
            for r in range(2, 425):
                ws.cell(row=r, column=1, value=marker)
            changed = True
            updated += 1
    if changed:
        wb.save(path)
        print(f'{cf}: updated', flush=True)

print(f'\nTotal sheets marker-filled: {updated}')
print(f'Still-empty sheets (non-2024, no pairs): {len(still_empty_non2024)}')
for e in still_empty_non2024:
    print(f'  {e[0]} | {e[1]} | {e[2]} {e[3]}')
