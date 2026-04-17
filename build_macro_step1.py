"""Macro Step 1: extract Macro-Debt_Data (2018gn) or Data-Input (2013gn) from each
source .xlsm and add as a new sheet in the existing _compiled/<country>.xlsx."""
import os, re, warnings, csv
warnings.filterwarnings('ignore')
from openpyxl import load_workbook

BASE = '/home/gabos/projects/external-debt-extraction/after-2018'
COMPILED_DIR = os.path.join(BASE, '_compiled')
SKIP_DIRS = {'Metadata', 'Completed_data_all_backup', '_compiled', '_compiled_step1_backup'}

INVALID_SHEET = re.compile(r'[\\/?*\[\]:]')

def base_id(fname):
    name = os.path.splitext(fname)[0]
    name = re.split(r'\s*,\s*', name)[0]
    name = re.split(r'[_\s\.]Sup', name, maxsplit=1, flags=re.IGNORECASE)[0]
    return name.strip()

def make_sheet_name(file_id, suffix, used):
    candidate = INVALID_SHEET.sub('_', f'{file_id}_{suffix}')
    if len(candidate) > 31:
        keep = 31 - len(suffix) - 1
        candidate = f'{file_id[:keep]}_{suffix}'
    if candidate not in used:
        return candidate
    i = 2
    while True:
        suf = f'_v{i}'
        truncd = candidate[:31 - len(suf)] + suf
        if truncd not in used:
            return truncd
        i += 1

countries = sorted(d for d in os.listdir(BASE)
                   if os.path.isdir(os.path.join(BASE, d)) and d not in SKIP_DIRS)
report = []
total_ok = total_skip = 0

for country in countries:
    cdir = os.path.join(BASE, country)
    files = sorted(f for f in os.listdir(cdir)
                   if f.lower().endswith(('.xlsm', '.xlsx')) and not f.startswith('~$'))
    if not files:
        continue
    cpath = os.path.join(COMPILED_DIR, f'{country}.xlsx')
    if not os.path.exists(cpath):
        print(f'WARN: no compiled file for {country}')
        continue
    wb = load_workbook(cpath)
    used = set(wb.sheetnames)
    n_done = n_skip = 0
    for fname in files:
        path = os.path.join(cdir, fname)
        try:
            src = load_workbook(path, data_only=True, read_only=True, keep_vba=False)
        except Exception as e:
            report.append((country, fname, 'ERROR_OPEN', str(e)[:80]))
            n_skip += 1
            continue
        sh = {s.lower(): s for s in src.sheetnames}
        if 'macro-debt_data' in sh:
            src_name = sh['macro-debt_data']; suffix = 'Macro_Debt_Data'
        elif 'data-input' in sh:
            src_name = sh['data-input']; suffix = 'Data_Input'
        else:
            report.append((country, fname, 'NO_MATCHING_SHEET', ''))
            n_skip += 1
            src.close()
            continue
        src_ws = src[src_name]
        file_id = base_id(fname)
        new_name = make_sheet_name(file_id, suffix, used)
        used.add(new_name)
        ws = wb.create_sheet(new_name)
        for row in src_ws.iter_rows(values_only=True):
            ws.append(row)
        src.close()
        report.append((country, fname, 'OK', new_name))
        n_done += 1
    # alphabetize sheets so new ones cluster nicely
    wb._sheets.sort(key=lambda s: s.title)
    wb.save(cpath)
    print(f'{country}: {n_done} added, {n_skip} skipped', flush=True)
    total_ok += n_done
    total_skip += n_skip

# write report
out = os.path.join(COMPILED_DIR, '_macro_step1_report.csv')
with open(out, 'w', newline='') as f:
    w = csv.writer(f); w.writerow(['country','source_file','status','detail'])
    w.writerows(report)
print(f'\nTotal sheets added: {total_ok}')
print(f'Skipped: {total_skip}')
print(f'Report: {out}')
