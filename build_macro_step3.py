"""Macro Step 3: delete 'Further Details' col, insert dsa.pdf as col A,
replace indicator column with MacroDebtData template col A values."""
import os, re, warnings
warnings.filterwarnings('ignore')
from openpyxl import load_workbook

COMPILED_DIR = '/home/gabos/projects/external-debt-extraction/after-2018/_compiled'
MACRO_TPL_PATH = '/home/gabos/projects/external-debt-extraction/DSA_Assumptions_MacroDebtData.xlsx'

# Load macro template col A
tpl_wb = load_workbook(MACRO_TPL_PATH, data_only=True)
tpl_ws = tpl_wb['Template_Draft']
MACRO_ROWS = tpl_ws.max_row  # 122
macro_col_a = [tpl_ws.cell(row=r, column=1).value for r in range(1, MACRO_ROWS + 1)]
tpl_wb.close()
print(f'Macro template col A loaded: {MACRO_ROWS} rows ({sum(1 for v in macro_col_a if v is not None)} non-empty)')

TARGET_SUFFIXES = ('_Macro_Debt_Data', '_Data_Input')

def sibling_sheet(cur_sheets, file_id):
    """Find the Ext_Debt_Data or Inp_Out_Debt sibling for a given file_id."""
    for suf in ('_Ext_Debt_Data', '_Inp_Out_Debt'):
        candidate = f'{file_id}{suf}'
        if candidate in cur_sheets:
            return candidate
    # try prefix match (Yemen-style truncations)
    for s in cur_sheets:
        if s.endswith('_Ext_Debt_Data') or s.endswith('_Inp_Out_Debt'):
            fid = re.sub(r'_(Ext_Debt_Data|Inp_Out_Debt)(?:_v\d+)?$', '', s)
            if fid.startswith(file_id) or file_id.startswith(fid):
                return s
    return None

total = 0; warned = []
for fname in sorted(os.listdir(COMPILED_DIR)):
    if not fname.endswith('.xlsx') or fname.startswith('_') or fname.startswith('~$'):
        continue
    path = os.path.join(COMPILED_DIR, fname)
    wb = load_workbook(path)
    sheets = wb.sheetnames[:]
    for sname in sheets:
        if not sname.endswith(TARGET_SUFFIXES): continue
        ws = wb[sname]
        # derive file_id by stripping macro/data suffix
        fid = re.sub(r'_(Macro_Debt_Data|Data_Input)(?:_v\d+)?$', '', sname)
        sib = sibling_sheet(sheets, fid)
        pdf_val = None
        if sib:
            pdf_val = wb[sib].cell(row=2, column=1).value
        else:
            warned.append(f'{fname}/{sname}: no sibling -> dsa.pdf blank')

        # Step 3a: delete col J (Further Details)
        ws.delete_cols(10, 1)

        # Step 3b: insert new col A and fill with dsa.pdf
        ws.insert_cols(1, 1)
        ws.cell(row=1, column=1, value='dsa.pdf')
        if pdf_val is not None:
            # fill all rows that have ANY content (or just 2..max_row) — use range up to max_row
            for r in range(2, ws.max_row + 1):
                ws.cell(row=r, column=1, value=pdf_val)

        # Step 3c: replace col I (indicator) with MacroDebtData col A
        # col I is column 9 (after dsa.pdf shift, the old col H 'indicator' moved to col I)
        # clear existing col I first up to 424 rows (where old template indicators lived)
        max_clear = max(424, MACRO_ROWS)
        for r in range(1, max_clear + 1):
            ws.cell(row=r, column=9, value=None)
        # write macro template col A
        for ri, v in enumerate(macro_col_a):
            if v is not None:
                ws.cell(row=ri + 1, column=9, value=v)

        total += 1
    wb.save(path)
    print(f'{fname}: saved', flush=True)

print(f'\nTotal macro sheets processed: {total}')
if warned:
    print('Warnings:')
    for w in warned[:20]: print(' ', w)
