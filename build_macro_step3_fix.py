"""Fix step 3 bug: properly clear col I beyond macro template and reapply macro col A.
Also re-integrate Haiti HTI_EBS.23.4 Ext_Debt_Data (was missing), and re-fill its
macro sheet's dsa.pdf column."""
import os, warnings, re
warnings.filterwarnings('ignore')
from openpyxl import load_workbook

COMPILED_DIR = '/home/gabos/projects/external-debt-extraction/after-2018/_compiled'
MACRO_TPL = '/home/gabos/projects/external-debt-extraction/DSA_Assumptions_MacroDebtData.xlsx'
TPL = '/home/gabos/projects/external-debt-extraction/DSA_Assumptions_ExternalDebtData_DomesticDebtData.xlsx'

# Load macro col A
w = load_workbook(MACRO_TPL, data_only=True)
mws = w['Template_Draft']
MACRO_ROWS = mws.max_row
macro_col_a = [mws.cell(row=r, column=1).value for r in range(1, MACRO_ROWS+1)]
w.close()

# 1) Re-integrate Haiti HTI_EBS.23.4 Ext_Debt_Data
SRC = '/home/gabos/projects/external-debt-extraction/after-2018/Haiti_HTI/HTI_EBS.23.4, Sup. 1.xlsm'
HCP = os.path.join(COMPILED_DIR, 'Haiti_HTI.xlsx')

w = load_workbook(TPL, data_only=True)
tpl_ws = w['TEMPLATE']
tpl_cells = [[tpl_ws.cell(row=r,column=c).value for c in range(1,13)] for r in range(1,425)]
w.close()

sw = load_workbook(SRC, data_only=True, keep_vba=False)
src_ws = sw['Ext_Debt_Data']
src_data = [[src_ws.cell(row=r,column=c).value for c in range(1, src_ws.max_column+1)]
            for r in range(1, src_ws.max_row+1)]
first_year = sw['Input 1 - Basics']['C18'].value
sw.close()

hwb = load_workbook(HCP)
SHEET = 'HTI_EBS.23.4_Ext_Debt_Data'
if SHEET in hwb.sheetnames:
    del hwb[SHEET]
ws = hwb.create_sheet(SHEET)
# paste template A1:L424
for r in range(424):
    for c in range(12):
        v = tpl_cells[r][c]
        if v is not None:
            ws.cell(row=r+1, column=c+1, value=v)
# source data at I1
for ri, row in enumerate(src_data):
    for ci, val in enumerate(row):
        if val is None: continue
        ws.cell(row=1+ri, column=9+ci, value=val)
# I1:L1 headers
ws.cell(row=1, column=9, value='description')
ws.cell(row=1, column=10, value='trs1')
ws.cell(row=1, column=11, value='trs2')
ws.cell(row=1, column=12, value='trs3')
# metadata rows 2:424
for r in range(2, 425):
    ws.cell(row=r, column=2, value='EBS.23.4, Sup. 1')
    ws.cell(row=r, column=3, value='Haiti')
    ws.cell(row=r, column=4, value='HTI')
    ws.cell(row=r, column=6, value='2018gn')
    ws.cell(row=r, column=7, value=int(first_year))
# Now apply the same col-A insertion this file had from step 5:
# the existing Ext_Debt_Data sheets already have col A = dsa.pdf
# For this newly-created sheet, insert col A and set to matching PDF (Haiti 2023 → check pairs)
# Easier: match with Haiti's other 2023 sheets — from pairs, Haiti 2023 = Haiti_2023_May (or whatever)
# Check existing Haiti 2023 sheet to reuse pattern
existing_2023 = [s for s in hwb.sheetnames if 'HTI_EBS.22' in s or 'HTI_EBS.23' in s or 'HTI_SM' in s or 'HTI_EBS.20' in s]
# I'll rely on the loaded pairs below for the macro step too, but HTI_EBS.23.4 published around 2023; check pair
# Manual mapping: Haiti pairs for 2023 — look them up now
import sys
from collections import defaultdict
pw = load_workbook('/home/gabos/projects/external-debt-extraction/dsa_parent_pairs.xlsx', read_only=True)
pairs = defaultdict(lambda: defaultdict(list))
for row in pw.active.iter_rows(min_row=2, values_only=True):
    fn,country,_,year,month = row[0],row[1],row[2],row[3],row[4]
    if fn and country and year and month:
        try: pairs[country][int(year)].append((int(month),fn))
        except: pass
pw.close()
for c in pairs:
    for y in pairs[c]: pairs[c][y].sort()
hti_2023 = pairs.get('Haiti', {}).get(2023, [])
# Haiti existing Ext 2023: only this new one (23.4). Check what doc num ordering says
# All Haiti 2023 ext sheets: find them now
all_hti_ext_2023 = sorted([s for s in hwb.sheetnames if s.startswith('HTI_') and '.23.' in s and s.endswith('_Ext_Debt_Data')])
# Our new sheet is HTI_EBS.23.4_Ext_Debt_Data, doc=4. Sort all_hti_ext_2023 by doc:
def doc_of(s):
    m = re.match(r'^HTI_\w+\.\d{2}\.(\d+)', s)
    return int(m.group(1)) if m else 0
all_hti_ext_2023.sort(key=doc_of)
# Pair up with hti_2023 sorted by month
pdf_for_new = None
for i, sh in enumerate(all_hti_ext_2023):
    if sh == SHEET and i < len(hti_2023):
        pdf_for_new = hti_2023[i][1]; break
# Shift col A for new sheet: insert col A
ws.insert_cols(1, 1)
ws.cell(row=1, column=1, value='dsa.pdf')
if pdf_for_new:
    for r in range(2, 425):
        ws.cell(row=r, column=1, value=pdf_for_new)
hwb._sheets.sort(key=lambda s: s.title)
hwb.save(HCP)
print(f'Haiti re-integrated: {SHEET} | dsa.pdf={pdf_for_new}')

# 2) Fix all macro sheets: properly clear col I and rewrite macro col A
# Also for Haiti HTI_EBS.23.4_Macro_Debt_Data, refill its dsa.pdf col using the sibling we just created
total = 0
for fname in sorted(os.listdir(COMPILED_DIR)):
    if not fname.endswith('.xlsx') or fname.startswith('_') or fname.startswith('~$'): continue
    path = os.path.join(COMPILED_DIR, fname)
    wb = load_workbook(path)
    sheets = wb.sheetnames[:]
    for sname in sheets:
        if not sname.endswith(('_Macro_Debt_Data','_Data_Input')): continue
        ws = wb[sname]
        # Clear col I rows 1..424 properly
        for r in range(1, 425):
            ws.cell(row=r, column=9).value = None
        # Write macro col A rows 1..122
        for ri, v in enumerate(macro_col_a):
            if v is not None:
                ws.cell(row=ri+1, column=9, value=v)
        total += 1
        # Haiti special: refill dsa.pdf if blank
        if fname == 'Haiti_HTI.xlsx' and sname == 'HTI_EBS.23.4_Macro_Debt_Data':
            if ws['A2'].value is None:
                sib = 'HTI_EBS.23.4_Ext_Debt_Data'
                if sib in wb.sheetnames:
                    pdf_val = wb[sib]['A2'].value
                    if pdf_val:
                        for r in range(2, ws.max_row+1):
                            ws.cell(row=r, column=1, value=pdf_val)
                        print(f'Haiti macro dsa.pdf filled: {pdf_val}')
    wb.save(path)
    print(f'{fname}: saved', flush=True)

print(f'\nFixed macro sheets: {total}')
