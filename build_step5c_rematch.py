"""Step 5c: unified re-match of dsa.pdf using all rules, overwriting col A."""
import os, re, warnings
from collections import defaultdict
warnings.filterwarnings('ignore')
from openpyxl import load_workbook

COMPILED_DIR = '/home/gabos/projects/external-debt-extraction/after-2018/_compiled'
PAIRS_PATH = '/home/gabos/projects/external-debt-extraction/dsa_parent_pairs.xlsx'

NAME_MAP = {
    'Congo DR': 'Congo, Dem. Rep',
    'Congo Republic of': 'Congo, Rep',
    "Cote d'Ivoire": "Côte d'Ivoire",
    'Gambia': 'Gambia, The',
    'Micronesia': 'Micronesia, Federated States of',
    'St. Vincent': 'St. Vincent and the Grenadines',
    'Yemen': 'Yemen, Rep',
}
TYPO_FIX = {'Afghanisatn': 'Afghanistan', "Cote d'lvoire": "Cote d'Ivoire", 'Tazania': 'Tanzania'}

# Load pairs
wb = load_workbook(PAIRS_PATH, read_only=True)
pairs = defaultdict(lambda: defaultdict(list))
for row in wb.active.iter_rows(min_row=2, values_only=True):
    fn, country, _, year, month = row[0], row[1], row[2], row[3], row[4]
    if fn and country and year and month:
        try: pairs[country][int(year)].append((int(month), fn))
        except: pass
wb.close()
for c in pairs:
    for y in pairs[c]: pairs[c][y].sort()

SHEET_RE = re.compile(r'^([A-Z]{3})_(EBS|SM)\.(\d{2})\.(\d+)')

# Gather entries per (file, country_pairs_name, year)
groups = defaultdict(list)  # (cf, cpairs, year) -> [(doc, sheet)]
for cf in sorted(os.listdir(COMPILED_DIR)):
    if not cf.endswith('.xlsx') or cf.startswith('_') or cf.startswith('~$'): continue
    folder = os.path.splitext(cf)[0]
    cname = TYPO_FIX.get(folder.rsplit('_', 1)[0].strip(), folder.rsplit('_', 1)[0].strip())
    cpairs = NAME_MAP.get(cname, cname)
    wb = load_workbook(os.path.join(COMPILED_DIR, cf), read_only=True)
    for s in wb.sheetnames:
        if s == 'EMPTY': continue
        m = SHEET_RE.match(s)
        if not m: continue
        yy = int(m.group(3))
        year = 2000 + yy if yy < 50 else 1900 + yy
        groups[(cf, cpairs, year)].append((int(m.group(4)), s))
    wb.close()

# Build assignments
assignments = defaultdict(dict)
counts = defaultdict(int)
for (cf, cpairs, year), ours in groups.items():
    ours.sort()
    plist = pairs.get(cpairs, {}).get(year, [])
    if not plist:
        if year == 2024: marker = 'no_pdf_2024'
        elif cpairs not in pairs: marker = 'no_country_in_pairs'
        else: marker = 'no_year_in_pairs'
        for doc, s in ours:
            assignments[cf][s] = marker
            counts[marker] += 1
        continue
    n = min(len(ours), len(plist))
    for i in range(n):
        assignments[cf][ours[i][1]] = plist[i][1]
        counts['MATCHED'] += 1
    for i in range(n, len(ours)):
        assignments[cf][ours[i][1]] = 'more_excel_than_pdf'
        counts['more_excel_than_pdf'] += 1

# Apply
saved = 0
for cf, smap in assignments.items():
    path = os.path.join(COMPILED_DIR, cf)
    wb = load_workbook(path)
    for sname, val in smap.items():
        ws = wb[sname]
        for r in range(2, 425):
            ws.cell(row=r, column=1, value=val)
    wb.save(path)
    saved += 1
    print(f'{cf}: saved', flush=True)

print(f'\nFiles saved: {saved}')
for k,v in sorted(counts.items()): print(f'  {k}: {v}')
