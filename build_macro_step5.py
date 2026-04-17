"""Macro Step 5: fill C (dsa.vintage), D (country.name), E (country.code)
for all macro sheets. Uses source filename and country folder."""
import os, re, warnings
warnings.filterwarnings('ignore')
from openpyxl import load_workbook

BASE = '/home/gabos/projects/external-debt-extraction/after-2018'
COMPILED_DIR = os.path.join(BASE, '_compiled')
SKIP = {'Metadata','Completed_data_all_backup','_compiled','_compiled_step1_backup','_compiled_macro_step1_backup'}
FILL_ROWS = range(2, 425)
COUNTRY_NAME_FIX = {
    'Afghanisatn': 'Afghanistan',
    "Cote d'lvoire": "Cote d'Ivoire",
    'Tazania': 'Tanzania',
}

def base_id(fname):
    name = os.path.splitext(fname)[0]
    name = re.split(r'\s*,\s*', name)[0]
    name = re.split(r'[_\s\.]Sup', name, maxsplit=1, flags=re.IGNORECASE)[0]
    return name.strip()

def vintage_from(fname):
    stem = os.path.splitext(fname)[0]
    m = re.match(r'^[A-Za-z]{3}\s*_\s*(.+)$', stem)
    return m.group(1).strip() if m else stem.strip()

def split_folder(folder):
    if '_' in folder:
        name, code = folder.rsplit('_', 1)
        name = COUNTRY_NAME_FIX.get(name.strip(), name.strip())
        return name, code.strip()
    return folder, ''

# country -> {file_id -> filename}
country_files = {}
for d in sorted(os.listdir(BASE)):
    if d in SKIP or not os.path.isdir(os.path.join(BASE, d)): continue
    m = {}
    for f in os.listdir(os.path.join(BASE, d)):
        if not f.lower().endswith(('.xlsm','.xlsx')) or f.startswith('~$'): continue
        m[base_id(f)] = f
    country_files[d] = m

total = 0; miss = 0
for cf in sorted(os.listdir(COMPILED_DIR)):
    if not cf.endswith('.xlsx') or cf.startswith('_') or cf.startswith('~$'): continue
    folder = os.path.splitext(cf)[0]
    cname, ccode = split_folder(folder)
    path = os.path.join(COMPILED_DIR, cf)
    wb = load_workbook(path)
    file_map = country_files.get(folder, {})
    changed = False
    for sname in wb.sheetnames:
        if not sname.endswith(('_Macro_Debt_Data', '_Data_Input')): continue
        ws = wb[sname]
        fid = re.sub(r'_(Macro_Debt_Data|Data_Input)(?:_v\d+)?$', '', sname)
        src = file_map.get(fid)
        if not src:
            # prefix match (Yemen-style truncation)
            cands = [fn for k, fn in file_map.items() if k.startswith(fid) or fid.startswith(k)]
            src = cands[0] if cands else None
        if not src:
            miss += 1; continue
        vintage = vintage_from(src)
        for r in FILL_ROWS:
            ws.cell(row=r, column=3, value=vintage)   # C dsa.vintage
            ws.cell(row=r, column=4, value=cname)     # D country.name
            ws.cell(row=r, column=5, value=ccode)     # E country.code
        changed = True
        total += 1
    if changed:
        wb.save(path)
        print(f'{folder}: {cname} / {ccode}', flush=True)

print(f'\nTotal: {total}, missing: {miss}')
