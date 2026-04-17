"""Macro Step 4: reset col I indicator to start from dod.ext.dbt.all at I2;
fill col G (methodo.gn) with 2018gn/2013gn."""
import os, warnings
warnings.filterwarnings('ignore')
from openpyxl import load_workbook

COMPILED_DIR = '/home/gabos/projects/external-debt-extraction/after-2018/_compiled'
MACRO_TPL = '/home/gabos/projects/external-debt-extraction/DSA_Assumptions_MacroDebtData.xlsx'

# Load macro col A and extract from dod.ext.dbt.all onwards
w = load_workbook(MACRO_TPL, data_only=True)
mws = w['Template_Draft']
MACRO_ROWS = mws.max_row
macro_col_a = [mws.cell(row=r, column=1).value for r in range(1, MACRO_ROWS + 1)]
w.close()

# find index of 'dod.ext.dbt.all'
start_idx = macro_col_a.index('dod.ext.dbt.all')
indicator_block = macro_col_a[start_idx:]  # from dod.ext.dbt.all to end
print(f'Indicator block from row {start_idx+1} ({macro_col_a[start_idx]!r}) '
      f'to row {MACRO_ROWS} — {len(indicator_block)} rows ({sum(1 for v in indicator_block if v)} non-empty)')

FILL_ROWS = range(2, 425)  # same as other metadata cols

total = 0
for fname in sorted(os.listdir(COMPILED_DIR)):
    if not fname.endswith('.xlsx') or fname.startswith('_') or fname.startswith('~$'):
        continue
    path = os.path.join(COMPILED_DIR, fname)
    wb = load_workbook(path)
    for sname in wb.sheetnames:
        if sname.endswith('_Macro_Debt_Data'):
            methodo = '2018gn'
        elif sname.endswith('_Data_Input'):
            methodo = '2013gn'
        else:
            continue
        ws = wb[sname]
        # Clear col I fully and rewrite
        max_clear = max(ws.max_row, 424)
        for r in range(1, max_clear + 1):
            ws.cell(row=r, column=9).value = None
        ws.cell(row=1, column=9, value='indicator')
        for ri, v in enumerate(indicator_block):
            if v is not None:
                ws.cell(row=2 + ri, column=9, value=v)
        # Fill col G (methodo.gn)
        for r in FILL_ROWS:
            ws.cell(row=r, column=7, value=methodo)
        total += 1
    wb.save(path)
    print(f'{fname}: saved', flush=True)

print(f'\nTotal macro sheets updated: {total}')
