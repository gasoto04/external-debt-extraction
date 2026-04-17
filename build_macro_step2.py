"""Macro Step 2: in each *_Macro_Debt_Data / *_Data_Input sheet, paste TEMPLATE
A1:L424 and place the source data at I1 with hard-set I1:L1 headers.

ONLY touches sheets ending in _Macro_Debt_Data or _Data_Input. Existing
_Ext_Debt_Data / _Inp_Out_Debt sheets are not modified.
"""
import os, warnings, shutil
warnings.filterwarnings('ignore')
from openpyxl import load_workbook

TPL_PATH = '/home/gabos/projects/external-debt-extraction/DSA_Assumptions_ExternalDebtData_DomesticDebtData.xlsx'
COMPILED_DIR = '/home/gabos/projects/external-debt-extraction/after-2018/_compiled'
BACKUP = '/home/gabos/projects/external-debt-extraction/after-2018/_compiled_macro_step1_backup'

print('Loading TEMPLATE A:L...')
tpl_wb = load_workbook(TPL_PATH, data_only=True)
tpl_ws = tpl_wb['TEMPLATE']
TPL_ROWS, TPL_COLS = 424, 12
template = [[tpl_ws.cell(row=r, column=c).value for c in range(1, TPL_COLS + 1)]
            for r in range(1, TPL_ROWS + 1)]
tpl_wb.close()

# backup
if not os.path.exists(BACKUP):
    print(f'Backup -> {BACKUP}')
    shutil.copytree(COMPILED_DIR, BACKUP)
else:
    print(f'Backup already exists: {BACKUP}')

target_suffixes = ('_Macro_Debt_Data', '_Data_Input')
total_sheets = 0
for fname in sorted(os.listdir(COMPILED_DIR)):
    if not fname.endswith('.xlsx') or fname.startswith('_') or fname.startswith('~$'):
        continue
    path = os.path.join(COMPILED_DIR, fname)
    wb = load_workbook(path)
    n = 0
    for sname in list(wb.sheetnames):
        if not sname.endswith(target_suffixes):
            continue
        ws = wb[sname]
        # snapshot source data
        src_data = [row for row in ws.iter_rows(values_only=True)]
        idx = wb.sheetnames.index(sname)
        del wb[sname]
        new_ws = wb.create_sheet(sname, idx)
        # paste template
        for r in range(TPL_ROWS):
            for c in range(TPL_COLS):
                v = template[r][c]
                if v is not None:
                    new_ws.cell(row=r + 1, column=c + 1, value=v)
        # paste source data at I1
        for ri, row in enumerate(src_data):
            for ci, val in enumerate(row):
                if val is None: continue
                new_ws.cell(row=1 + ri, column=9 + ci, value=val)
        # hard-set I1:L1 headers
        new_ws.cell(row=1, column=9, value='description')
        new_ws.cell(row=1, column=10, value='trs1')
        new_ws.cell(row=1, column=11, value='trs2')
        new_ws.cell(row=1, column=12, value='trs3')
        n += 1
    wb.save(path)
    print(f'{fname}: {n} macro sheets updated', flush=True)
    total_sheets += n

print(f'\nTotal macro sheets updated: {total_sheets}')
print(f'Pre-step2 backup: {BACKUP}')
