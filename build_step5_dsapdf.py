"""Step 5: shift all data +1 column right, insert new col A = dsa.pdf + PDF filename."""
import os, re, warnings
from collections import defaultdict
warnings.filterwarnings('ignore')
from openpyxl import load_workbook

COMPILED_DIR = '/home/gabos/projects/external-debt-extraction/after-2018/_compiled'
PAIRS_PATH = '/home/gabos/projects/external-debt-extraction/dsa_parent_pairs.xlsx'

NAME_MAP = {
    'Congo DR': 'Congo, Dem. Rep',
    'Congo Republic of': 'Congo, Rep',
    'Cote d\'Ivoire': 'Côte d\'Ivoire',
    'Gambia': 'Gambia, The',
    'Micronesia': 'Micronesia, Federated States of',
    'St. Vincent': 'St. Vincent and the Grenadines',
    'Yemen': 'Yemen, Rep',
}
TYPO_FIX = {'Afghanisatn': 'Afghanistan', "Cote d'lvoire": "Cote d'Ivoire", 'Tazania': 'Tanzania'}

# Load pairs: country_in_pairs -> {year -> [(month, filename)]}
print('Loading pairs...')
wb = load_workbook(PAIRS_PATH, read_only=True)
ws = wb.active
pairs = defaultdict(lambda: defaultdict(list))
for row in ws.iter_rows(min_row=2, values_only=True):
    fn, country, ifs, year, month = row[0], row[1], row[2], row[3], row[4]
    if fn and country and year and month:
        try:
            pairs[country][int(year)].append((int(month), fn))
        except (TypeError, ValueError):
            pass
wb.close()
for c in pairs:
    for y in pairs[c]:
        pairs[c][y].sort()  # ascending by month

# Gather all (compiled_file, sheet, country_pairs_name, year, doc_num)
SHEET_RE = re.compile(r'^([A-Z]{3})_(EBS|SM)\.(\d{2})\.(\d+)')
entries = []
for cf in sorted(os.listdir(COMPILED_DIR)):
    if not cf.endswith('.xlsx') or cf.startswith('_') or cf.startswith('~$'): continue
    folder = os.path.splitext(cf)[0]
    cname_raw = folder.rsplit('_', 1)[0].strip()
    cname = TYPO_FIX.get(cname_raw, cname_raw)
    cname_pairs = NAME_MAP.get(cname, cname)
    wb = load_workbook(os.path.join(COMPILED_DIR, cf), read_only=True)
    for s in wb.sheetnames:
        if s == 'EMPTY': continue
        m = SHEET_RE.match(s)
        if not m: continue
        yy = int(m.group(3))
        year = 2000 + yy if yy < 50 else 1900 + yy
        doc = int(m.group(4))
        entries.append((cf, s, cname_pairs, year, doc))
    wb.close()

# Group by (compiled_file, country_pairs_name, year) and match
groups = defaultdict(list)
for cf, s, cname_pairs, year, doc in entries:
    groups[(cf, cname_pairs, year)].append((doc, s))

# Build {file -> {sheet -> pdf_filename}}
assignments = defaultdict(dict)
unmatched = []
for (cf, cname_pairs, year), ours in groups.items():
    ours.sort()  # ascending by doc_num
    pair_list = pairs.get(cname_pairs, {}).get(year, [])
    if not pair_list:
        for doc, s in ours:
            assignments[cf][s] = None
            unmatched.append((cf, s, cname_pairs, year, doc, 'NO_PAIRS_FOR_COUNTRY_YEAR'))
        continue
    # if counts differ, still match by index up to min(len)
    n = min(len(ours), len(pair_list))
    for i in range(n):
        doc, s = ours[i]
        month, pdf = pair_list[i]
        assignments[cf][s] = pdf
    # leftovers
    if len(ours) > n:
        for i in range(n, len(ours)):
            doc, s = ours[i]
            assignments[cf][s] = None
            unmatched.append((cf, s, cname_pairs, year, doc, f'MORE_SHEETS_THAN_PAIRS ({len(ours)}>{len(pair_list)})'))
    if len(pair_list) > n:
        extras = [pdf for _,pdf in pair_list[n:]]
        unmatched.append((cf, '(extra pairs)', cname_pairs, year, None, f'EXTRA_PAIRS: {extras}'))

print(f'\nGroups: {len(groups)}, total sheets to write: {sum(len(v) for v in assignments.values())}')
print(f'Unmatched/issues: {len(unmatched)}')

# Apply to each file: insert_cols(1) on every non-EMPTY sheet, then write dsa.pdf
total_saved = 0
for cf in sorted(assignments.keys()):
    path = os.path.join(COMPILED_DIR, cf)
    wb = load_workbook(path)
    for sname in wb.sheetnames:
        if sname == 'EMPTY': continue
        ws = wb[sname]
        ws.insert_cols(1)  # everything shifts +1 right
        ws.cell(row=1, column=1, value='dsa.pdf')
        pdf = assignments[cf].get(sname)
        if pdf:
            for r in range(2, 425):
                ws.cell(row=r, column=1, value=pdf)
    wb.save(path)
    total_saved += 1
    print(f'{cf}: saved', flush=True)

print(f'\nFiles saved: {total_saved}')
print('\nUnmatched details:')
for u in unmatched:
    print(' ', u)
