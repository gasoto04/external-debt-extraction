"""Macro Step 6: copy first.year.proj (col H) and issue.date (col F) from sibling."""
import os, re, warnings
warnings.filterwarnings('ignore')
from openpyxl import load_workbook

COMPILED_DIR = '/home/gabos/projects/external-debt-extraction/after-2018/_compiled'
FILL_ROWS = range(2, 425)

def sibling(cur_sheets, file_id):
    for suf in ('_Ext_Debt_Data', '_Inp_Out_Debt'):
        c = f'{file_id}{suf}'
        if c in cur_sheets: return c
    for s in cur_sheets:
        if s.endswith(('_Ext_Debt_Data','_Inp_Out_Debt')):
            fid = re.sub(r'_(Ext_Debt_Data|Inp_Out_Debt)(?:_v\d+)?$', '', s)
            if fid.startswith(file_id) or file_id.startswith(fid):
                return s
    return None

total = 0; miss = 0
for cf in sorted(os.listdir(COMPILED_DIR)):
    if not cf.endswith('.xlsx') or cf.startswith('_') or cf.startswith('~$'): continue
    path = os.path.join(COMPILED_DIR, cf)
    wb = load_workbook(path)
    sheets = wb.sheetnames[:]
    changed = False
    for sname in sheets:
        if not sname.endswith(('_Macro_Debt_Data','_Data_Input')): continue
        ws = wb[sname]
        fid = re.sub(r'_(Macro_Debt_Data|Data_Input)(?:_v\d+)?$', '', sname)
        sib = sibling(sheets, fid)
        if not sib:
            miss += 1; continue
        sib_ws = wb[sib]
        fy = sib_ws['H2'].value   # first.year.proj in sibling Ext
        idate = sib_ws['F2'].value  # issue.date
        idate_fmt = sib_ws['F2'].number_format
        for r in FILL_ROWS:
            if fy is not None:
                ws.cell(row=r, column=8, value=fy)     # H first.year.proj
            if idate is not None:
                c = ws.cell(row=r, column=6, value=idate)  # F issue.date
                c.number_format = idate_fmt
        changed = True
        total += 1
    if changed:
        wb.save(path)
        print(f'{cf}: saved', flush=True)

print(f'\nTotal: {total}, missing sibling: {miss}')
