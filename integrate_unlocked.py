"""Integrate previously-locked files into compiled country workbooks."""
import os, re, warnings
warnings.filterwarnings('ignore')
from openpyxl import load_workbook

COMPILED_DIR = '/home/gabos/projects/external-debt-extraction/after-2018/_compiled'
TPL_PATH = '/home/gabos/projects/external-debt-extraction/DSA_Assumptions_ExternalDebtData_DomesticDebtData.xlsx'

TARGETS = [
    ('Nicaragua_NIC', 'Nicaragua', 'NIC', 'NIC_SM.23.6, Sup. 2.xlsm'),
    ('Papua New Guinea_PNG', 'Papua New Guinea', 'PNG', 'PNG_EBS.20.109, Sup. 1.xlsm'),
    ('Zambia_ZMB', 'Zambia', 'ZMB', 'ZMB_EBS.22.73, Sup. 1.xlsm'),
    ('Zambia_ZMB', 'Zambia', 'ZMB', 'ZMB_EBS.23.146, Sup. 1.xlsm'),
    ('Zambia_ZMB', 'Zambia', 'ZMB', 'ZMB_EBS.23.77, Sup. 2.xlsm'),
    ('Zambia_ZMB', 'Zambia', 'ZMB', 'ZMB_SM.19.188, Sup. 2.xlsm'),
]

def base_id(fname):
    name = os.path.splitext(fname)[0]
    name = re.split(r'\s*,\s*', name)[0]
    name = re.split(r'[_\s\.]Sup', name, maxsplit=1, flags=re.IGNORECASE)[0]
    return name.strip()

def vintage_from(fname):
    stem = os.path.splitext(fname)[0]
    m = re.match(r'^[A-Za-z]{3}\s*_\s*(.+)$', stem)
    return m.group(1).strip() if m else stem.strip()

# Load template A1:L424
tpl_wb = load_workbook(TPL_PATH, data_only=True)
tpl_ws = tpl_wb['TEMPLATE']
tpl_cells = [[tpl_ws.cell(row=r, column=c).value for c in range(1, 13)] for r in range(1, 425)]
tpl_wb.close()

for folder, cname, ccode, fname in TARGETS:
    src_path = f'/home/gabos/projects/external-debt-extraction/after-2018/{folder}/{fname}'
    compiled_path = os.path.join(COMPILED_DIR, f'{folder}.xlsx')
    sheet_name = f'{base_id(fname)}_Ext_Debt_Data'
    vintage = vintage_from(fname)

    src_wb = load_workbook(src_path, data_only=True, keep_vba=False)
    src_ws = src_wb['Ext_Debt_Data']
    src_data = [[src_ws.cell(row=r, column=c).value for c in range(1, src_ws.max_column + 1)]
                for r in range(1, src_ws.max_row + 1)]
    sh = {s.lower(): s for s in src_wb.sheetnames}
    fy = None
    if 'input 1 - basics' in sh:
        fy = src_wb[sh['input 1 - basics']]['C18'].value
    elif 'output - submit' in sh:
        fy = src_wb[sh['output - submit']]['C9'].value
    src_wb.close()
    if not isinstance(fy, (int, float)):
        print(f'  WARN: no first_year for {fname}, using None'); fy = None

    wb = load_workbook(compiled_path)
    # remove EMPTY placeholder if present (Zambia case)
    if 'EMPTY' in wb.sheetnames:
        del wb['EMPTY']
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)

    # Template A:L
    for r in range(424):
        for c in range(12):
            v = tpl_cells[r][c]
            if v is not None:
                ws.cell(row=r + 1, column=c + 1, value=v)
    # Source data at I1
    for ri, row in enumerate(src_data):
        for ci, val in enumerate(row):
            if val is None: continue
            ws.cell(row=1 + ri, column=9 + ci, value=val)
    # I1:L1 headers
    ws.cell(row=1, column=9, value='description')
    ws.cell(row=1, column=10, value='trs1')
    ws.cell(row=1, column=11, value='trs2')
    ws.cell(row=1, column=12, value='trs3')
    # Metadata rows 2:424
    for r in range(2, 425):
        ws.cell(row=r, column=2, value=vintage)
        ws.cell(row=r, column=3, value=cname)
        ws.cell(row=r, column=4, value=ccode)
        ws.cell(row=r, column=6, value='2018gn')
        if fy is not None:
            ws.cell(row=r, column=7, value=int(fy))
    # sort sheets alphabetically
    wb._sheets.sort(key=lambda s: s.title)
    wb.save(compiled_path)
    print(f'{folder}: added {sheet_name} (first_year={fy})')
